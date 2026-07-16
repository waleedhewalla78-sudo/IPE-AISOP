#!/usr/bin/env python3
"""Generate Excel fixtures for Phases 3-5 strategy upload tests."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from openpyxl import Workbook
except ImportError:
    print("openpyxl required: uv pip install openpyxl", file=sys.stderr)
    raise

OUT = Path(__file__).resolve().parents[1] / "docs" / "qa" / "test-fixtures" / "phases3-5"


def _save(wb: Workbook, name: str) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    wb.save(path)
    return path


def product_master_valid() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "product_master"
    ws.append(["product_code", "name", "type", "uom", "name_ar"])
    rows = [
        ("FG-DT100", "Distribution Transformer 100kVA", "finished", "unit", "محول توزيع 100 ك.ف.أ"),
        ("FG-DT250", "Distribution Transformer 250kVA", "finished", "unit", "محول توزيع 250 ك.ف.أ"),
        ("FG-PT500", "Power Transformer 500kVA", "finished", "unit", "محول قوة 500 ك.ف.أ"),
        ("SM-COIL100", "Coil Assembly DT100", "semi", "unit", ""),
        ("SM-COIL250", "Coil Assembly DT250", "semi", "unit", ""),
        ("SM-COIL500", "Coil Assembly PT500", "semi", "unit", ""),
        ("SM-CORE", "Core Stack", "semi", "unit", ""),
        ("RM-CW25", "Copper Wire 25mm", "raw", "kg", ""),
        ("RM-SS", "Silicon Steel Sheet", "raw", "kg", ""),
        ("RM-INS", "Insulation Paper", "raw", "kg", ""),
        ("RM-OIL", "Transformer Oil", "raw", "litre", ""),
        ("RM-BOLT", "Mounting Bolts", "raw", "unit", ""),
    ]
    for row in rows:
        ws.append(list(row))
    return _save(wb, "product_master_valid.xlsx")


def product_master_missing_column() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "type", "uom"])
    ws.append(["FG-DT100", "finished", "unit"])
    return _save(wb, "product_master_missing_product_code.xlsx")


def product_master_bad_types() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["product_code", "name", "type", "uom", "unit_cost", "quantity"])
    ws.append(["FG-DT100", "Transformer", "finished", "unit", "expensive", "5"])
    ws.append(["FG-DT250", "Transformer 250", "finished", "unit", "1200", "-5"])
    ws.append(["RM-CW25", "Copper Wire", "raw", "kg", "45.5", "100"])
    return _save(wb, "product_master_bad_types.xlsx")


def product_master_duplicates() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["product_code", "name", "type", "uom"])
    ws.append(["FG-DT100", "First DT100", "finished", "unit"])
    ws.append(["FG-DT100", "Duplicate DT100", "finished", "unit"])
    ws.append(["FG-DT250", "DT250", "finished", "unit"])
    return _save(wb, "product_master_duplicates.xlsx")


def product_master_empty() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["product_code", "name", "type", "uom"])
    return _save(wb, "product_master_empty.xlsx")


def product_master_arabic() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["product_code", "name", "type", "uom", "name_ar"])
    ws.append(["FG-AR01", "محول توزيع 100 ك.ف.أ", "finished", "unit", "محول توزيع 100 ك.ف.أ"])
    return _save(wb, "product_master_arabic.xlsx")


def bom_orphan() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["product_code", "component_code", "quantity"])
    ws.append(["FG-DT999", "RM-CW25", "10"])
    ws.append(["FG-DT100", "RM-CW25", "5"])
    return _save(wb, "bom_orphan_parent.xlsx")


def sales_orders_smoke(n: int = 500) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "order_number",
            "customer_code",
            "product_code",
            "quantity",
            "order_date",
            "requested_delivery",
            "status",
        ]
    )
    for i in range(n):
        ws.append(
            [
                f"SO-PERF-{i:05d}",
                "CUST-SEC",
                "FG-DT100",
                "2",
                "2026-07-01",
                "2026-08-15",
                "confirmed",
            ]
        )
    return _save(wb, f"sales_orders_{n}_rows.xlsx")


def sales_orders_valid() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "order_number",
            "customer_code",
            "product_code",
            "quantity",
            "order_date",
            "requested_delivery",
            "status",
        ]
    )
    ws.append(["SO-2026-0200", "CUST-SEC", "FG-DT100", "5", "2026-07-10", "2026-08-10", "confirmed"])
    return _save(wb, "sales_orders_valid.xlsx")


def fake_pdf_as_xlsx() -> Path:
    path = OUT / "not_really.xlsx"
    OUT.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"%PDF-1.4 fake content renamed to xlsx")
    return path


def main() -> None:
    paths = [
        product_master_valid(),
        product_master_missing_column(),
        product_master_bad_types(),
        product_master_duplicates(),
        product_master_empty(),
        product_master_arabic(),
        bom_orphan(),
        sales_orders_valid(),
        sales_orders_smoke(500),
        fake_pdf_as_xlsx(),
    ]
    print(f"Generated {len(paths)} fixtures in {OUT}")
    for p in paths:
        print(f"  {p.name}")


if __name__ == "__main__":
    main()
