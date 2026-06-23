"""Parse and validate project plan Excel workbooks."""

from __future__ import annotations

import hashlib
import io
from typing import Any

from app.core.project_plan_schema import (
    ALLOWED_EXTENSIONS,
    ALLOWED_STATUSES,
    MAX_FILE_BYTES,
    REQUIRED_HEADERS,
    normalize_header,
)

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - dependency declared in pyproject
    load_workbook = None


class ProjectPlanParseError(Exception):
    def __init__(self, message: str, errors: list[str] | None = None):
        super().__init__(message)
        self.errors = errors or [message]

    def __str__(self) -> str:
        base = self.args[0] if self.args else "Validation failed"
        if len(self.errors) == 1 and self.errors[0] != base:
            return self.errors[0]
        if len(self.errors) == 1:
            return self.errors[0]
        if self.errors:
            preview = "; ".join(self.errors[:3])
            if len(self.errors) > 3:
                preview += f"; ... and {len(self.errors) - 3} more"
            return f"{base}: {preview}"
        return base


def _validate_extension(file_name: str) -> None:
    lower = file_name.lower()
    if not any(lower.endswith(ext) for ext in ALLOWED_EXTENSIONS):
        raise ProjectPlanParseError(
            f"Invalid file type. Upload an Excel workbook ({', '.join(ALLOWED_EXTENSIONS)})."
        )


def _coerce_float(value: Any, field: str, row_num: int) -> float:
    if value is None or value == "":
        raise ProjectPlanParseError(f"Row {row_num}: {field} is required")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ProjectPlanParseError(f"Row {row_num}: {field} must be a number") from exc
    return result


def _coerce_int(value: Any, field: str, row_num: int) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError) as exc:
        raise ProjectPlanParseError(f"Row {row_num}: {field} must be an integer") from exc


def parse_project_plan_excel(file_name: str, content: bytes) -> dict[str, Any]:
    """Parse workbook bytes into normalized plan payload."""
    if load_workbook is None:
        raise ProjectPlanParseError("Excel parser unavailable (openpyxl not installed)")

    _validate_extension(file_name)
    if len(content) == 0:
        raise ProjectPlanParseError("Uploaded file is empty")
    if len(content) > MAX_FILE_BYTES:
        raise ProjectPlanParseError(f"File exceeds maximum size of {MAX_FILE_BYTES // (1024 * 1024)} MB")

    try:
        workbook = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ProjectPlanParseError("Could not read Excel file — file may be corrupted") from exc

    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()

    if not rows:
        raise ProjectPlanParseError("Workbook is empty")

    header_row = rows[0]
    column_map: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        normalized = normalize_header(cell)
        if normalized and normalized not in column_map:
            column_map[normalized] = idx

    missing = [h for h in REQUIRED_HEADERS if h not in column_map]
    if missing:
        raise ProjectPlanParseError(
            f"Missing required columns: {', '.join(missing)}",
            [f"Missing column: {col}" for col in missing],
        )

    operations: list[dict[str, Any]] = []
    errors: list[str] = []
    plan_codes: set[str] = set()
    plan_names: set[str] = set()
    seen_keys: set[tuple[str, int]] = set()

    for row_idx, row in enumerate(rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        def cell(name: str) -> Any:
            col = column_map.get(name)
            if col is None or col >= len(row):
                return None
            return row[col]

        try:
            plan_code = str(cell("PLAN_CODE")).strip()
            plan_name = str(cell("PLAN_NAME")).strip()
            mo_id = str(cell("MO_ID")).strip()
            op_name = str(cell("OPERATION_NAME")).strip()
            wc_code = str(cell("WORK_CENTER_CODE")).strip()
            sequence = _coerce_int(cell("OPERATION_SEQUENCE"), "OPERATION_SEQUENCE", row_idx)
            start_hour = _coerce_float(cell("START_HOUR"), "START_HOUR", row_idx)
            duration_hours = _coerce_float(cell("DURATION_HOURS"), "DURATION_HOURS", row_idx)
        except ProjectPlanParseError as exc:
            errors.extend(exc.errors)
            continue

        if not plan_code:
            errors.append(f"Row {row_idx}: PLAN_CODE is required")
            continue
        if not plan_name:
            errors.append(f"Row {row_idx}: PLAN_NAME is required")
            continue
        if not mo_id:
            errors.append(f"Row {row_idx}: MO_ID is required")
            continue
        if not wc_code:
            errors.append(f"Row {row_idx}: WORK_CENTER_CODE is required")
            continue
        if duration_hours <= 0:
            errors.append(f"Row {row_idx}: DURATION_HOURS must be greater than 0")
            continue
        if start_hour < 0:
            errors.append(f"Row {row_idx}: START_HOUR must be >= 0")
            continue

        status_raw = cell("STATUS")
        status = str(status_raw).strip().lower() if status_raw not in (None, "") else "planned"
        if status not in ALLOWED_STATUSES:
            errors.append(
                f"Row {row_idx}: STATUS must be one of {', '.join(sorted(ALLOWED_STATUSES))}"
            )
            continue

        key = (mo_id, sequence)
        if key in seen_keys:
            errors.append(f"Row {row_idx}: duplicate MO_ID + OPERATION_SEQUENCE ({mo_id}, {sequence})")
            continue
        seen_keys.add(key)

        plan_codes.add(plan_code)
        plan_names.add(plan_name)

        notes = cell("NOTES")
        operations.append(
            {
                "plan_code": plan_code,
                "plan_name": plan_name,
                "mo_id": mo_id,
                "operation_sequence": sequence,
                "operation_name": op_name or f"Op {sequence}",
                "work_center_code": wc_code,
                "start_hour": round(start_hour, 4),
                "duration_hours": round(duration_hours, 4),
                "start_minute": round(start_hour * 60, 2),
                "duration_minutes": round(duration_hours * 60, 2),
                "end_minute": round((start_hour + duration_hours) * 60, 2),
                "status": status,
                "notes": str(notes).strip() if notes not in (None, "") else None,
            }
        )

    if errors:
        raise ProjectPlanParseError("Validation failed", errors)

    if not operations:
        raise ProjectPlanParseError("No operation rows found below the header row")

    if len(plan_codes) > 1:
        raise ProjectPlanParseError(
            "All rows must share the same PLAN_CODE",
            [f"Found multiple plan codes: {', '.join(sorted(plan_codes))}"],
        )
    if len(plan_names) > 1:
        raise ProjectPlanParseError(
            "All rows must share the same PLAN_NAME",
            [f"Found multiple plan names: {', '.join(sorted(plan_names))}"],
        )

    file_hash = hashlib.sha256(content).hexdigest()
    plan_code = next(iter(plan_codes))
    plan_name = next(iter(plan_names))

    return {
        "plan_code": plan_code,
        "plan_name": plan_name,
        "operations": operations,
        "row_count": len(operations),
        "file_sha256": file_hash,
        "file_size_bytes": len(content),
    }
